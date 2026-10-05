import csv
from collections.abc import Callable
from io import BytesIO
from typing import Final

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from analytics.dashboard import Table

XLSX_CONTENT_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_COLUMN_WIDTH: Final = 60


def table_to_csv(table: Table) -> HttpResponse:
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = _attachment(table, "csv")
    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow(table.headers)
    writer.writerows(table.rows)
    return response


def table_to_xlsx(table: Table) -> HttpResponse:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = table.name.capitalize()
    sheet.append(table.headers)
    for row in table.rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for index, column in enumerate(sheet.iter_cols(values_only=True), start=1):
        width = max(len(str(value)) for value in column if value is not None)
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = _attachment(table, "xlsx")
    return response


EXPORTERS: Final[dict[str, Callable[[Table], HttpResponse]]] = {
    "csv": table_to_csv,
    "xlsx": table_to_xlsx,
}


def _attachment(table: Table, extension: str) -> str:
    return f'attachment; filename="jobpulse-{table.name}.{extension}"'
